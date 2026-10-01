from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pandas as pd
from docx import Document as DocxDocument
from langchain_core.documents import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".xlsx", ".csv"}


def load_documents(filename: str, file_bytes: bytes) -> list[Document]:
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type. Supported types: {supported}")

    if extension == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))
        return [
            Document(
                page_content=page.extract_text() or "",
                metadata={"source": filename, "page": page_number},
            )
            for page_number, page in enumerate(reader.pages, start=1)
        ]

    if extension == ".docx":
        document = DocxDocument(BytesIO(file_bytes))
        text = "\n".join(
            paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
        )
        return [Document(page_content=text, metadata={"source": filename})]

    if extension == ".txt":
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = file_bytes.decode("latin-1", errors="replace")
        return [Document(page_content=text, metadata={"source": filename})]

    if extension == ".xlsx":
        sheets = pd.read_excel(BytesIO(file_bytes), sheet_name=None)
        text = "\n\n".join(
            f"Sheet: {sheet_name}\n{dataframe.to_csv(index=False)}"
            for sheet_name, dataframe in sheets.items()
        )
        return [Document(page_content=text, metadata={"source": filename})]

    dataframe = pd.read_csv(BytesIO(file_bytes))
    return [
        Document(
            page_content=dataframe.to_csv(index=False),
            metadata={"source": filename},
        )
    ]